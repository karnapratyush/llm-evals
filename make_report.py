"""
Turn model_results.csv (written by run_models.py) into an Excel report: eval_report.xlsx

Sheets:
  Scoreboard        - one row per model: overall + per-difficulty accuracy, latency, failure reasons
  By Skill          - accuracy per SQL skill tag (join, window, anti-join, ...) per model
  Head to Head      - model vs model: questions the row model got right and the column model got wrong
  Question Matrix   - one row per question, a pass/fail cell per model, and how many models solved it
  Hardest Questions - the 10 questions the fewest models solved, with the gold SQL to study
  SQL Side by Side  - gold SQL + gold result next to every model's SQL, result and verdict
  All Results       - the raw rows from model_results.csv (filterable)

run_models.py calls this automatically after every model; you can also rerun it on its own.
Several result files can be merged into one report - when the same (model, question) appears more
than once, the LATER file wins (e.g. a re-run of one question replaces the original answer):
  uv run python make_report.py
  uv run python make_report.py --input results/a.csv results/b.csv --output results/combined.xlsx
"""

import argparse

import pandas as pd
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import REPORT_PATH, RESULTS_PATH

DIFFICULTIES = ["easy", "medium", "hard", "brutal"]
PASS, FAIL = "PASS", "FAIL"

GREEN = PatternFill("solid", start_color="C6EFCE")
RED = PatternFill("solid", start_color="FFC7CE")
HEADER = PatternFill("solid", start_color="1F4E78")


# ---------- sheet builders (each returns a DataFrame) ----------
def failure_kind(reason):
    """Collapse detailed reasons ('row_count 3 vs gold 1') into a few buckets."""
    if reason == "match":
        return "match"
    for prefix, kind in [("row_count", "wrong_row_count"), ("col_count", "wrong_col_count")]:
        if str(reason).startswith(prefix):
            return kind
    return reason


def scoreboard(df):
    g = df.groupby("model")
    board = pd.DataFrame({
        "overall_%": g["correct"].mean() * 100,
        "correct": g["correct"].sum().astype(int),
        "total": g["correct"].count(),
    })
    by_diff = df.pivot_table(index="model", columns="difficulty", values="correct", aggfunc="mean") * 100
    for d in DIFFICULTIES:
        if d in by_diff.columns:
            board[f"{d}_%"] = by_diff[d]
    board["avg_latency_s"] = g["latency_s"].mean()
    if "cost_usd" in df.columns:   # older result files have no usage columns
        board["total_cost_usd"] = g["cost_usd"].sum().round(4)
        board["avg_completion_tokens"] = g["completion_tokens"].mean()
        board["avg_reasoning_tokens"] = g["reasoning_tokens"].mean()
    fails = (df[~df["correct"]].assign(kind=lambda x: x["reason"].map(failure_kind))
               .pivot_table(index="model", columns="kind", values="id", aggfunc="count", fill_value=0))
    board = board.join(fails.add_prefix("fail: "), how="left").fillna(0)
    board = board.sort_values("overall_%", ascending=False).round(1)
    if "total_cost_usd" in board.columns:   # re-add unrounded: costs are fractions of a dollar
        board["total_cost_usd"] = g["cost_usd"].sum().round(4)
    board = board.reset_index()
    board.insert(0, "rank", range(1, len(board) + 1))
    return board


def by_skill(df, order):
    skills = df.assign(skill=df["categories"].str.split(",")).explode("skill")
    table = skills.pivot_table(index="skill", columns="model", values="correct", aggfunc="mean") * 100
    table = table[order]
    table.insert(0, "questions", skills.groupby("skill")["id"].nunique())
    return table.sort_values("questions", ascending=False).round(1).reset_index()


def head_to_head(df, order):
    wide = df.pivot_table(index="id", columns="model", values="correct", aggfunc="first")[order]
    models = order
    out = pd.DataFrame(index=models, columns=models, dtype="Int64")
    for a in models:
        for b in models:
            out.loc[a, b] = pd.NA if a == b else int((wide[a] & ~wide[b]).sum())
    out.index.name = "row model right, column model wrong"
    return out.reset_index()


def question_matrix(df, order):
    info = df.groupby("id")[["difficulty", "categories", "question"]].first()
    wide = df.pivot_table(index="id", columns="model", values="correct", aggfunc="first")[order]
    solved = wide.sum(axis=1).astype(int)
    marks = wide.apply(lambda col: col.map({True: PASS, False: FAIL}))
    out = info.join(marks)
    out["models_correct"] = solved
    out["models_total"] = wide.shape[1]
    return out.reset_index()


def hardest(df, order, n=10):
    qm = question_matrix(df, order)
    gold = df.groupby("id")[["gold_sql", "gold_result"]].first()
    hard = qm[qm["models_correct"] < qm["models_total"]].set_index("id").join(gold)
    return hard.sort_values(["models_correct", "id"]).head(n).reset_index()


def side_by_side(df, order):
    base = df.groupby("id")[["difficulty", "question", "gold_sql", "gold_result"]].first()
    parts = [base]
    for model in order:
        m = df[df["model"] == model].set_index("id")
        parts.append(pd.DataFrame({
            f"{model} | verdict": m["correct"].map({True: PASS, False: FAIL}),
            f"{model} | reason": m["reason"],
            f"{model} | sql": m["model_sql"],
            f"{model} | result": m["model_result"],
        }))
    return pd.concat(parts, axis=1).reset_index()


# ---------- formatting ----------
def style_sheet(ws, df, widths=None, wrap_cols=(), pct_cols=(), passfail=True):
    """Header style, freeze panes, filters, widths, wrapping and conditional colours."""
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    last_row = ws.max_row

    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        width = (widths or {}).get(col)
        if width is None:
            sample = df[col].head(50).map(lambda v: len(str(v))).max() if len(df) else 0
            width = min(max(len(str(col)), sample) + 2, 40)
        ws.column_dimensions[letter].width = width

        if col in wrap_cols or any(col.endswith(w) for w in wrap_cols):
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].alignment = Alignment(wrap_text=True, vertical="top")
        if col in pct_cols and last_row >= 2:
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{last_row}",
                ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                               mid_type="num", mid_value=50, mid_color="FFEB84",
                               end_type="num", end_value=100, end_color="63BE7B"))

    if passfail and last_row >= 2:   # an empty sheet has no data range to colour
        area = f"A2:{get_column_letter(len(df.columns))}{last_row}"
        ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[f'"{PASS}"'], fill=GREEN))
        ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[f'"{FAIL}"'], fill=RED))


def load_results(paths):
    """Read one or more results CSVs; later files override earlier ones for the same (model, id)."""
    paths = [paths] if isinstance(paths, (str, bytes)) or not hasattr(paths, "__iter__") else list(paths)
    df = pd.concat([pd.read_csv(p) for p in paths], ignore_index=True)
    df = df.drop_duplicates(["model", "id"], keep="last").sort_values(["model", "id"], ignore_index=True)
    df["correct"] = df["correct"].astype(str).str.upper().eq("TRUE")
    return df


def build_report(results_paths=RESULTS_PATH, report_path=REPORT_PATH):
    df = load_results(results_paths)
    board = scoreboard(df)
    models = list(board["model"])   # best model first, on every sheet

    sheets = {
        "Scoreboard": board,
        "By Skill": by_skill(df, models),
        "Head to Head": head_to_head(df, models),
        "Question Matrix": question_matrix(df, models),
        "Hardest Questions": hardest(df, models),
        "SQL Side by Side": side_by_side(df, models),
        "All Results": df,
    }

    with pd.ExcelWriter(report_path, engine="openpyxl") as xw:
        for name, sheet in sheets.items():
            sheet.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            cols = list(sheet.columns)
            if name == "Scoreboard":
                style_sheet(ws, sheet, pct_cols=[c for c in cols if c.endswith("_%")])
            elif name == "By Skill":
                style_sheet(ws, sheet, pct_cols=models)
            elif name == "Head to Head":
                style_sheet(ws, sheet, widths={cols[0]: 36})
            elif name == "Question Matrix":
                style_sheet(ws, sheet, widths={"question": 70}, wrap_cols=["question"])
            elif name == "Hardest Questions":
                style_sheet(ws, sheet, widths={"question": 60, "gold_sql": 70, "gold_result": 40},
                            wrap_cols=["question", "gold_sql", "gold_result"])
            elif name == "SQL Side by Side":
                widths = {"question": 50, "gold_sql": 60, "gold_result": 35}
                widths.update({c: 60 for c in cols if c.endswith("| sql")})
                widths.update({c: 35 for c in cols if c.endswith("| result")})
                style_sheet(ws, sheet, widths=widths,
                            wrap_cols=["question", "gold_sql", "gold_result", "| sql", "| result"])
            else:
                style_sheet(ws, sheet, widths={"question": 50, "gold_sql": 50, "model_sql": 50,
                                               "gold_result": 30, "model_result": 30})

    print(f"Report saved -> {report_path}  ({len(models)} models, {df['id'].nunique()} questions)")
    return report_path


if __name__ == "__main__":
    p = argparse.ArgumentParser(description="Build the Excel eval report from one or more results CSVs")
    p.add_argument("--input", nargs="+", default=[RESULTS_PATH], help="results CSV(s); later files win")
    p.add_argument("--output", default=REPORT_PATH)
    a = p.parse_args()
    build_report(a.input, a.output)
